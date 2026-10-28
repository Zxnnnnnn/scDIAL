"""Unified training and compact two-file result export."""
from __future__ import annotations

import copy
import json
import os
from pathlib import Path
import tempfile

import h5py
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from threadpoolctl import threadpool_limits

from .data import resolve_single_omics_file, read_var_names
from .evaluation import evaluate_run
from .markers import MARKER_COLUMNS
from .training import main as train_main

CLUSTER_METRICS = ("ACC", "NMI", "Purity", "ARI", "Fscore", "Precision", "Recall")
EXPORT_MARKER_COLUMNS = (
    "cluster", "gene", "gene_index", "rank", "score", "dominant_fa",
    "fa_gene_weight", "source_gene_index",
)


def json_value(value):
    if isinstance(value, dict):
        return {str(k): json_value(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_value(v) for v in value]
    if isinstance(value, np.ndarray):
        return json_value(value.tolist())
    if isinstance(value, np.generic):
        return json_value(value.item())
    if isinstance(value, float) and not np.isfinite(value):
        return None
    if isinstance(value, Path):
        return str(value)
    return value


def write_json(path, value):
    temporary = path.with_suffix(".json.tmp")
    with temporary.open("w", encoding="utf-8") as stream:
        json.dump(json_value(value), stream, ensure_ascii=False, allow_nan=False, indent=2)
        stream.write("\n")
    os.replace(str(temporary), str(path))


def clustering_summary(evaluated):
    """The original seven clustering means in percentage units."""
    return {name: round(float(evaluated["clustering_metrics"][name]["mean"]) * 100, 2)
            for name in CLUSTER_METRICS}


def compact_results(evaluated, parameters):
    """Configuration and clustering only; omit selected epoch and FA summaries."""
    return json_value({
        "parameters": {k: v for k, v in parameters.items()
                       if k not in ("best_epoch", "best_by_metric", "data_root", "output_dir", "save_file")},
        "clustering_metrics": clustering_summary(evaluated),
    })


def result_directory(dataset, output_root):
    """One stable directory per dataset, reused by subsequent runs."""
    label = Path(str(dataset)).name
    if label in ("", ".", ".."):
        raise ValueError("Dataset name must identify an output subdirectory.")
    return Path(output_root).resolve() / label


def write_marker_workbook(path, markers):
    """Scale a copy for presentation; preserve the raw ranking and training data."""
    if len(markers) and not set(EXPORT_MARKER_COLUMNS).issubset(markers.columns):
        raise ValueError("Marker export requires all eight output columns.")
    presentation = markers.reindex(columns=EXPORT_MARKER_COLUMNS).copy()
    presentation["score"] = presentation["score"] * 10.0
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "marker_genes"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(list(EXPORT_MARKER_COLUMNS))
    for row in presentation.itertuples(index=False, name=None):
        sheet.append([None if pd.isna(value) else json_value(value) for value in row])
    widths = (11, 20, 15, 10, 16, 17, 22, 23)
    for column, width in zip("ABCDEFGH", widths):
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244A68")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 26
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=11)
            cell.alignment = Alignment(horizontal="left" if cell.column == 2 else "right",
                                       vertical="center")
            cell.number_format = "0.000000" if cell.column in (5, 7) else ("@" if cell.column == 2 else "0")
        # Gene symbols are literal text, including any leading formula character.
        row[1].data_type = "s"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(str(path))
    workbook.close()


def save_results(directory, markers, evaluated, parameters):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    # Prepare both outputs before replacing the previous successful results.
    with tempfile.TemporaryDirectory(prefix=".export_", dir=str(directory)) as temporary:
        staged = Path(temporary)
        write_marker_workbook(staged / "marker_genes.xlsx", markers)
        write_json(staged / "results.json", compact_results(evaluated, parameters))
        for name in ("marker_genes.xlsx", "results.json"):
            os.replace(str(staged / name), str(directory / name))
    legacy_csv = directory / "marker_genes.csv"
    if legacy_csv.is_file():
        legacy_csv.unlink()


def run_dataset(dataset, args):
    directory = result_directory(dataset, args.output_dir)
    runs, marker_frames = {}, {}
    roles = [("control", 0.0), ("new_method", args.lambda_marker)] if args.run_control else [("new_method", args.lambda_marker)]
    for role, weight in roles:
        effective_args = copy.deepcopy(args)
        effective_args.lambda_marker = weight
        with threadpool_limits(limits=args.threads):
            training, artifacts = train_main(dataset, effective_args, collect_artifacts=True)
            evaluated, marker_frame = evaluate_run(training, artifacts, effective_args)
        runs[role], marker_frames[role] = evaluated, marker_frame
        del artifacts, training

    markers = marker_frames["new_method"]
    if markers is None:
        markers = pd.DataFrame(columns=MARKER_COLUMNS)
    elif args.data_type == "single":
        path = resolve_single_omics_file(dataset, args.data_root)
        with h5py.File(str(path), "r") as source:
            n_features = len(source["var_names"])
        full_names = read_var_names(path, n_features)
        first = {}
        for index, gene in enumerate(full_names):
            first.setdefault(gene, index)
        markers["source_gene_index"] = markers["gene"].map(first)
    parameters = vars(args).copy()
    parameters["datasets"] = dataset
    parameters.update(runs["new_method"].get("marker_input_configuration", {}))
    save_results(directory, markers, runs["new_method"], parameters)
    print("Results: {}".format(directory), flush=True)
    return directory
